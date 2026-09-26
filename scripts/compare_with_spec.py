"""観測した vt_code を、地理院が公開している「注記分類コード・地物種別コード一覧」と突き合わせる。

仕様書(Excel)は実行のたびに地理院のサイトから取得し、リポジトリには保存しない。
観測は data/observations/vector_tiles.csv(observe_vector_tiles.py の出力)を使う。

出力:
  data/observations/spec_comparison.csv       コードごとの仕様上のズーム範囲と観測したズーム
  data/observations/spec_comparison.meta.json 仕様書の版(Last-Modified)と集計
"""

from __future__ import annotations

import argparse
import csv
import io
import json
from collections import defaultdict
from pathlib import Path

import openpyxl

from _gsi import PoliteSession, utc_now

ROOT = Path(__file__).resolve().parent.parent
SPEC_URL = "https://maps.gsi.go.jp/help/pdf/vector/optbv_featurecodes.xlsx"
ZOOMS = list(range(4, 18))
MARKS = {"○", "〇"}
# 面のレイヤは、スタイルで使っているものだけを見る
AREA_LAYERS = {"WA", "AdmArea"}


def zooms_of(cells) -> list[int]:
    return [z for z, v in zip(ZOOMS, cells) if v in MARKS]


def read_annotation_codes(wb) -> dict[tuple[str, str], tuple[str, list[int]]]:
    spec = {}
    for row in list(wb["注記分類コード"].iter_rows(values_only=True))[1:]:
        if isinstance(row[1], int):
            spec[("Anno", str(row[1]))] = (row[2], zooms_of(row[3:17]))
    return spec


def read_feature_codes(wb) -> dict[tuple[str, str], tuple[str, list[int]]]:
    """地物種別コードのシートは、上2桁と下2桁が別の列に分かれ、空欄は上の行を引き継ぐ。"""
    spec = {}
    layer = klass = head = label = None
    for row in list(wb["地物種別コード"].iter_rows(values_only=True))[1:]:
        if row[0]:
            layer, klass, head, label = row[0], row[1], None, None
        if layer not in AREA_LAYERS:
            continue
        if row[2] not in (None, "-"):
            head = str(row[2])
        if row[4] not in (None, "-"):
            label = str(row[4]).replace("\n", "")
        detail = f" / {row[5]}" if row[5] else ""
        # 「-」はコードを持たない地物(AdmArea の z4–z7)。
        code = "" if row[2] == "-" else f"{head}{row[3]}" if row[3] else None
        if code is None:
            continue
        key = (layer, code)
        name = f"{label}{detail}" if label else klass
        zs = zooms_of(row[6:20])
        if key in spec:  # 同じコードが縮尺帯ごとに別の行で出てくる
            spec[key] = (spec[key][0], sorted(set(spec[key][1]) | set(zs)))
        else:
            spec[key] = (name, zs)
    return spec


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--observations", type=Path, default=ROOT / "data" / "observations" / "vector_tiles.csv")
    parser.add_argument("--out", type=Path, default=ROOT / "data" / "observations")
    args = parser.parse_args()

    res = PoliteSession().get(SPEC_URL)
    res.raise_for_status()
    wb = openpyxl.load_workbook(io.BytesIO(res.content), read_only=True)
    spec = read_annotation_codes(wb) | read_feature_codes(wb)

    observed: dict[tuple[str, str], set[int]] = defaultdict(set)
    examples: dict[tuple[str, str], str] = {}
    with open(args.observations, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            key = (row["layer"], row["vt_code"])
            if row["layer"] == "Anno" and not row["vt_code"]:
                continue
            if row["layer"] != "Anno" and row["layer"] not in AREA_LAYERS:
                continue
            observed[key].add(int(row["zoom"]))
            if row["examples"] and key not in examples:
                examples[key] = row["examples"]

    rows = []
    for key in sorted(observed, key=lambda k: (k[0], int(k[1]) if k[1] else -1)):
        name, spec_zooms = spec.get(key, ("", []))
        outside = sorted(observed[key] - set(spec_zooms))
        status = "not_in_spec" if key not in spec else "outside_spec_zoom" if outside else "consistent"
        rows.append({
            "layer": key[0],
            "vt_code": key[1],
            "spec_name": name,
            "spec_zooms": " ".join(map(str, spec_zooms)),
            "observed_zooms": " ".join(map(str, sorted(observed[key]))),
            "outside_zooms": " ".join(map(str, outside)),
            "status": status,
            "examples": examples.get(key, ""),
        })

    args.out.mkdir(parents=True, exist_ok=True)
    with open(args.out / "spec_comparison.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    counts = defaultdict(int)
    for row in rows:
        counts[row["status"]] += 1
    meta = {
        "checked_at": utc_now(),
        "spec": SPEC_URL,
        "spec_last_modified": res.headers.get("Last-Modified"),
        "observations": str(args.observations.relative_to(ROOT)),
        "layers": ["Anno", *sorted(AREA_LAYERS)],
        "counts": dict(counts),
        "note": "Observed zooms come from a sample of tiles; a code not observed at a zoom may still exist there.",
    }
    (args.out / "spec_comparison.meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(meta["counts"], ensure_ascii=False))
    for row in rows:
        if row["status"] != "consistent":
            print(row)


if __name__ == "__main__":
    main()
